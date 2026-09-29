"""Extract the six position sheets of the skills workbook into build/positions.json.

Usage: python scripts/extract_positions.py [workbook.xlsx] [--out build/positions.json]
The workbook is not committed; this script only runs locally.
"""
import json
import os
import sys
import warnings

warnings.filterwarnings("ignore")
import openpyxl  # noqa: E402

# sheet -> (code, first_row, last_row, name_col, desc_col, tasks_col, level_col,
#           instructions_cell, level_ref_cols, level_ref_first_row,
#           objectives_cols, objectives_example_rows, job_title_cell)
LAYOUT = {
    "COO":  ("coo",  8, 19, "B", "C", None, "D", "B6", ("F", "G"), 6, ("F", "G"), (13, 16), "C4"),
    "CISO": ("ciso", 8, 19, "A", "B", None, "C", "A6", ("E", "F"), 6, ("E", "F"), (12, 15), "B4"),
    "CTO":  ("cto",  7, 17, "A", "B", None, "C", "A5", ("E", "F"), 5, ("E", "F"), (12, 15), "B3"),
    "CMO":  ("cmo",  8, 17, "B", "C", "D",  "E", "B6", ("G", "H"), 6, ("G", "H"), (13, 16), "C4"),
    "CSO":  ("cso",  8, 18, "B", "C", "D",  "E", "B6", ("G", "H"), 6, ("G", "H"), (13, 16), "C4"),
    "CFO":  ("cfo",  8, 16, "B", None, "C", "D", "B6", ("F", "G"), 6, ("F", "G"), (13, 16), "C4"),
}


def _s(v):
    if v is None:
        return None
    v = str(v).replace("﻿", "").replace(" ", "\n").strip()
    return v or None


def extract(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for sheet, (code, r0, r1, nc, dc, tc, lc, instr, lvl_cols, lvl_r0, obj_cols, obj_rows, title) in LAYOUT.items():
        ws = wb[sheet]
        cols = ["name"] + (["description"] if dc else []) + (["tasks"] if tc else []) + ["level"]
        skills = []
        for i, r in enumerate(range(r0, r1 + 1), start=1):
            full = _s(ws[f"{nc}{r}"].value)
            lvl = ws[f"{lc}{r}"].value
            skills.append({
                "n": i,
                "name_fr": full.split("\n")[0].strip() if full else None,
                "name_full_fr": full,
                "description_fr": _s(ws[f"{dc}{r}"].value) if dc else None,
                "tasks_fr": _s(ws[f"{tc}{r}"].value) if tc else None,
                "level": int(lvl) if isinstance(lvl, (int, float)) else None,
                "cells": f"{nc}{r}:{lc}{r}",
            })
        levels = {str(k): _s(ws[f"{lvl_cols[1]}{lvl_r0 + k - 1}"].value) for k in (1, 2, 3)}
        objs = [[_s(ws[f"{obj_cols[0]}{r}"].value), _s(ws[f"{obj_cols[1]}{r}"].value)]
                for r in range(obj_rows[0], obj_rows[1] + 1)]
        # The Department and Position manager labels sit in the instructions column,
        # three rows and one row above the instructions row; their values are one column right.
        val_col = chr(ord(instr[0]) + 1)
        instr_row = int(instr[1:])
        dep = _s(ws[f"{val_col}{instr_row - 3}"].value)
        man = _s(ws[f"{val_col}{instr_row - 1}"].value)
        out[code] = {
            "sheet": sheet,
            "job_title_fr": _s(ws[title].value),
            "department": dep,
            "manager": man,
            "instructions": _s(ws[instr].value),
            "levels": levels,
            "columns": cols,
            "skills": skills,
            "objectives_examples": objs,
            "objectives_cells": f"{obj_cols[0]}{obj_rows[0] - 2}:{obj_cols[1]}{obj_rows[1]}",
        }
    return out


# Cells that hold people's names in the workbook's team sheets. Only used to build the
# git-ignored forbidden-names list; the names themselves are never written to tracked files.
PEOPLE_CELLS = {
    "Skill Matrix": ["I5"] + [f"F{r}" for r in range(27, 67)],
    "Training Log": ["I5"] + [f"F{r}" for r in range(16, 61)],
    "Team Member Review": ["C6", "W6", "AS9"],
}


def forbidden_names(path):
    """Return lowercase name tokens (3+ chars) found in the workbook's people cells, sorted and unique."""
    wb = openpyxl.load_workbook(path, data_only=True)
    tokens = set()
    for sheet, cells in PEOPLE_CELLS.items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for c in cells:
            v = _s(ws[c].value)
            if not v or v.startswith("#"):
                continue
            for tok in v.replace("-", " ").split():
                tok = tok.strip(".,;:()").lower()
                if len(tok) >= 3:
                    tokens.add(tok)
    return sorted(tokens)


def main(argv):
    wb = argv[1] if len(argv) > 1 and not argv[1].startswith("--") else "the source skills workbook"
    outp = "build/positions.json"
    if "--out" in argv:
        outp = argv[argv.index("--out") + 1]
    if not os.path.exists(wb):
        print(f"workbook not found: {wb}")
        return 2
    if "--names" in argv:
        os.makedirs("build", exist_ok=True)
        toks = forbidden_names(wb)
        with open("build/forbidden-names.txt", "w", encoding="utf-8") as f:
            f.write("\n".join(toks) + "\n")
        print(f"wrote build/forbidden-names.txt: {len(toks)} tokens")
        return 0
    data = extract(wb)
    os.makedirs(os.path.dirname(outp), exist_ok=True)
    with open(outp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"wrote {outp}: " + ", ".join(f"{k}={len(v['skills'])}" for k, v in data.items()))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
